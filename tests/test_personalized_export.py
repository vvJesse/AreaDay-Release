from __future__ import annotations

import importlib.util
import io
import sys
import unittest
from email.message import Message
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / "scripts"
sys.path.insert(0, str(SCRIPTS))

import personalized_export as export  # noqa: E402


SPEC = importlib.util.spec_from_file_location(
    "researchramp_personalized_export_tests",
    ROOT / "app" / "server.py",
)
APP = importlib.util.module_from_spec(SPEC)
assert SPEC.loader is not None
sys.modules[SPEC.name] = APP
SPEC.loader.exec_module(APP)


CARD = {
    "lemma": "specie",
    "display_form": "specie",
    "part_of_speech": "noun",
    "meaning_en": "n. a coin of small value",
    "meaning_zh": "n. 硬币",
    "context": "The specie was counted twice in one paragraph.",
    "source_title": "A source paper",
    "source_url": "https://example.invalid/paper",
    "source_paper_id": "W123",
    "total_count": 36,
    "document_count": 12,
}

LABELS = [label for label, _field in export.WORKBOOK_COLUMNS]


def row(**overrides: object) -> dict[str, str]:
    """One calibration row, as the exported tab-separated file holds it."""

    base = {
        "lemma": "specie",
        "part_of_speech": "NOUN",
        "probability_known": "0.0123",
        "classification": "likely_unknown",
        "total_count": "36",
        "document_count": "12",
        "zipf": "3.0",
        "frequency_prior_probability": "0.10",
        "cefr_level": "B2",
        "cefr_adjustment": "0.00",
        "exam_tags": "",
        "exam_adjustment": "0.00",
        "education_adjustment": "0.00",
        "direct_response": "",
        "importance_tier": "A",
        "important_boundary_protected": "false",
        "selected_threshold": "0.9",
    }
    base.update({key: str(value) for key, value in overrides.items()})
    return base


def only_card(card: dict | None = CARD):
    def lookup(lemma: str, part_of_speech: str = "") -> dict | None:
        if card is None or str(lemma).casefold() != str(card["lemma"]).casefold():
            return None
        return dict(card)

    return lookup


def no_card(lemma: str, part_of_speech: str = "") -> None:
    return None


class PersonalizedSheetTests(unittest.TestCase):
    def build(self, rows: list[dict[str, str]], lookup=None):
        payload = export.build_personalized_workbook(rows, lookup or only_card())
        return load_workbook(BytesIO(payload))

    def cells(self, sheet, line: int) -> dict[str, object]:
        return {label: sheet.cell(row=line, column=index).value
                for index, label in enumerate(LABELS, start=1)}

    def test_the_sheet_shows_the_fifteen_columns_a_reader_reads(self) -> None:
        workbook = self.build([row()])
        self.assertEqual(workbook.sheetnames, [export.SHEET_SELECTED, export.SHEET_OTHER])
        self.assertEqual([cell.value for cell in workbook[export.SHEET_SELECTED][1]], LABELS)
        self.assertEqual(
            LABELS,
            ["词", "词性", "领域英文释义", "领域中文释义", "论文原句", "论文标题", "论文链接",
             "论文标识", "收录状态", "优先级", "预测掌握度", "30 词回答", "覆盖论文数",
             "出现次数", "本次阈值"],
        )

    def test_the_reviewed_card_supplies_the_meaning_the_sentence_and_the_paper(self) -> None:
        cells = self.cells(self.build([row()])[export.SHEET_SELECTED], 2)
        self.assertEqual(cells["词"], "specie")
        self.assertEqual(cells["词性"], "noun")
        self.assertEqual(cells["领域英文释义"], CARD["meaning_en"])
        self.assertEqual(cells["领域中文释义"], CARD["meaning_zh"])
        self.assertEqual(cells["论文原句"], CARD["context"])
        self.assertEqual(cells["论文标题"], CARD["source_title"])
        self.assertEqual(cells["论文链接"], CARD["source_url"])
        self.assertEqual(cells["论文标识"], CARD["source_paper_id"])
        self.assertEqual(cells["覆盖论文数"], 12)
        self.assertEqual(cells["出现次数"], 36)

    def test_a_word_the_corpus_has_no_card_for_keeps_the_row_it_came_from(self) -> None:
        cells = self.cells(self.build([row()], no_card)[export.SHEET_SELECTED], 2)
        self.assertEqual(cells["词"], "specie")
        self.assertEqual(cells["词性"], "NOUN")
        self.assertIsNone(cells["领域中文释义"])
        self.assertIsNone(cells["论文链接"])
        self.assertEqual(cells["覆盖论文数"], 12)
        self.assertEqual(cells["优先级"], "A")

    def test_a_word_the_run_counts_as_known_goes_to_the_other_candidates(self) -> None:
        workbook = self.build(
            [row(lemma="specie"), row(lemma="charlie", classification="likely_known")]
        )
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_SELECTED]["A"]][1:], ["specie"]
        )
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_OTHER]["A"]][1:], ["charlie"]
        )

    def test_the_words_to_learn_come_first_then_tier_then_certainty(self) -> None:
        workbook = self.build(
            [
                row(lemma="delta", importance_tier="D", probability_known="0.90"),
                row(lemma="alpha", importance_tier="A", probability_known="0.80"),
                row(lemma="bravo", importance_tier="A", probability_known="0.20"),
                row(lemma="charlie", classification="likely_known", importance_tier="A"),
            ]
        )
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_SELECTED]["A"]][1:],
            ["bravo", "alpha", "delta"],
        )
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_OTHER]["A"]][1:], ["charlie"]
        )

    def test_words_certain_in_the_same_way_are_ordered_by_their_corpus_support(self) -> None:
        workbook = self.build(
            [
                row(lemma="bravo", document_count="2", total_count="9"),
                row(lemma="charlie", document_count="5", total_count="1"),
                row(lemma="alpha", document_count="5", total_count="1"),
            ]
        )
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_SELECTED]["A"]][1:],
            ["alpha", "charlie", "bravo"],
        )

    def test_the_reader_sees_a_percentage_and_the_words_own_answer(self) -> None:
        workbook = self.build(
            [
                row(lemma="specie", probability_known="0.873", direct_response="known"),
                row(lemma="charlie", probability_known="0.0625", direct_response=""),
            ]
        )
        sheet = workbook[export.SHEET_SELECTED]
        uncertain, likely_known_word = self.cells(sheet, 2), self.cells(sheet, 3)
        self.assertEqual(uncertain["词"], "charlie")
        self.assertEqual(uncertain["预测掌握度"], "6.3%")
        self.assertEqual(uncertain["30 词回答"], "未抽到")
        self.assertEqual(likely_known_word["预测掌握度"], "87.3%")
        self.assertEqual(likely_known_word["30 词回答"], "known")
        self.assertEqual(likely_known_word["本次阈值"], "90%")

    def test_the_inclusion_column_names_what_happens_to_the_word(self) -> None:
        workbook = self.build(
            [
                row(lemma="alpha", classification="uncertain"),
                row(
                    lemma="bravo",
                    classification="important_boundary",
                    important_boundary_protected="true",
                ),
                row(lemma="charlie", classification="likely_known"),
                row(
                    lemma="delta",
                    classification="likely_known",
                    important_boundary_protected="true",
                ),
            ]
        )
        sheet = workbook[export.SHEET_SELECTED]
        self.assertEqual(
            [self.cells(sheet, line)["收录状态"] for line in (2, 3)],
            ["建议加入", "重要词保护后加入"],
        )
        other = workbook[export.SHEET_OTHER]
        self.assertEqual(
            [self.cells(other, line)["词"] for line in (2, 3)], ["charlie", "delta"]
        )
        self.assertEqual(
            [self.cells(other, line)["收录状态"] for line in (2, 3)],
            ["暂不加入", "暂不加入"],
        )

    def test_a_sentence_with_a_control_character_stays_readable(self) -> None:
        card = {**CARD, "context": "water ranged between 13.7\x01C and 16.9\x01C."}
        workbook = self.build([row()], only_card(card))
        self.assertEqual(
            self.cells(workbook[export.SHEET_SELECTED], 2)["论文原句"],
            "water ranged between 13.7_x0001_C and 16.9_x0001_C.",
        )

    def test_both_sheets_carry_the_header_the_filter_and_the_widths(self) -> None:
        workbook = self.build([row()])
        for name in workbook.sheetnames:
            sheet = workbook[name]
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, f"A1:O{sheet.max_row}")
            self.assertEqual(sheet.column_dimensions["A"].width, 14.83203125)
            self.assertEqual(sheet.column_dimensions["E"].width, 60.83203125)
            self.assertEqual(sheet.column_dimensions["F"].width, 36.83203125)
            self.assertEqual(sheet.column_dimensions["G"].width, 46.83203125)

    def test_a_run_without_rows_has_no_sheet_to_hand_over(self) -> None:
        with self.assertRaises(ValueError):
            export.build_personalized_workbook([], only_card())

    def test_the_exported_file_reads_back_into_the_same_rows(self) -> None:
        source = (
            "lemma\tpart_of_speech\tprobability_known\tclassification\ttotal_count\t"
            "document_count\tdirect_response\timportance_tier\t"
            "important_boundary_protected\tselected_threshold\n"
            "specie\tNOUN\t0.0123\tlikely_unknown\t36\t12\t\tA\tfalse\t0.9\n"
        )
        rows = export.read_export_rows(source)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["lemma"], "specie")
        self.assertEqual(rows[0]["selected_threshold"], "0.9")
        workbook = self.build(rows)
        self.assertEqual(self.cells(workbook[export.SHEET_SELECTED], 2)["词"], "specie")


class DownloadRouteTests(unittest.TestCase):
    """The download button hands over the spreadsheet, under its own name."""

    def setUp(self) -> None:
        self.export = (
            "lemma\tpart_of_speech\tprobability_known\tclassification\ttotal_count\t"
            "document_count\tdirect_response\timportance_tier\t"
            "important_boundary_protected\tselected_threshold\n"
            "specie\tNOUN\t0.0123\tlikely_unknown\t36\t12\t\tA\tfalse\t0.9\n"
        )
        session = SimpleNamespace(persisted_export_tsv=lambda: self.export)
        context = SimpleNamespace(
            domain_id="alpha", session=session, continuous_store=None
        )
        runtime = SimpleNamespace(
            contexts={"alpha": context},
            context=lambda domain_id=None: context,
            catalog_card=lambda context, lemma, part_of_speech="": (
                APP.AppRuntime.catalog_card(context, lemma, part_of_speech)
            ),
        )

        class Handler(APP.AppHandler):
            def __init__(handler_self, path: str) -> None:
                handler_self.command = "GET"
                handler_self.path = path
                handler_self.rfile = io.BytesIO(b"")
                handler_self.wfile = io.BytesIO()
                handler_self.headers = Message()
                handler_self.response_status = None
                handler_self.sent_headers = {}

            def send_response(handler_self, code, message=None) -> None:
                handler_self.response_status = code

            def send_header(handler_self, keyword: str, value: str) -> None:
                handler_self.sent_headers[keyword] = value

            def end_headers(handler_self) -> None:
                return

        Handler.runtime = runtime
        Handler.static_dir = ROOT / "app" / "static"
        self.handler = Handler("/api/export.xlsx?domain_id=alpha")

    def test_the_download_is_a_named_spreadsheet(self) -> None:
        self.handler.do_GET()
        self.assertEqual(self.handler.response_status, 200)
        self.assertEqual(
            self.handler.sent_headers["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertEqual(
            self.handler.sent_headers["Content-Disposition"],
            'attachment; filename="personalized-vocabulary.xlsx"',
        )
        workbook = load_workbook(BytesIO(self.handler.wfile.getvalue()))
        self.assertEqual(workbook.sheetnames, [export.SHEET_SELECTED, export.SHEET_OTHER])
        self.assertEqual(
            [cell.value for cell in workbook[export.SHEET_SELECTED][1]], LABELS
        )
        self.assertEqual(
            workbook[export.SHEET_SELECTED].cell(row=2, column=1).value, "specie"
        )

    def test_the_page_asks_for_the_spreadsheet_by_name(self) -> None:
        page = (ROOT / "app" / "static" / "index.html").read_text(encoding="utf-8")
        script = (ROOT / "app" / "static" / "app.js").read_text(encoding="utf-8")
        self.assertIn('href="/api/export.xlsx"', page)
        self.assertIn('download="personalized-vocabulary.xlsx"', page)
        self.assertIn('domainUrl("/api/export.xlsx", currentDomainId)', script)
        self.assertNotIn("/api/export.tsv", script)


if __name__ == "__main__":
    unittest.main()
